# excel_writer.py
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from datetime import datetime
import os
from config import EXCEL_REPORT_DIR

def write_event_to_excel(user_id, username, action, timestamp_str):
    timestamp = datetime.strptime(timestamp_str, "%Y-%m-%d %H:%M:%S")
    month_name = timestamp.strftime('%B')
    date_str = timestamp.strftime('%Y-%m-%d')
    time_str = timestamp.strftime('%H:%M:%S')

    file_path = os.path.join(EXCEL_REPORT_DIR, f"report_{user_id}_2025_by_months.xlsx")

    if not os.path.exists(EXCEL_REPORT_DIR):
        os.makedirs(EXCEL_REPORT_DIR)

    if os.path.exists(file_path):
        wb = load_workbook(file_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if month_name not in wb.sheetnames:
        ws = wb.create_sheet(title=month_name)
        ws.append(["Дата", "Начало", "Обед-выход", "Обед-возврат", "Конец"])
    else:
        ws = wb[month_name]

    row_found = False
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == date_str:
            row_found = True
            if action == "Пришел на работу":
                ws.cell(row=row, column=2).value = time_str
            elif action == "Вышел на обед":
                ws.cell(row=row, column=3).value = time_str
            elif action == "Вернулся с обеда":
                ws.cell(row=row, column=4).value = time_str
            elif action == "Ушел с работы":
                ws.cell(row=row, column=5).value = time_str
            break

    if not row_found:
        new_row = [date_str, '', '', '', '']
        if action == "Пришел на работу":
            new_row[1] = time_str
        elif action == "Вышел на обед":
            new_row[2] = time_str
        elif action == "Вернулся с обеда":
            new_row[3] = time_str
        elif action == "Ушел с работы":
            new_row[4] = time_str
        ws.append(new_row)

    for col in ws.columns:
        for cell in col:
            cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))
            cell.alignment = Alignment(horizontal='center')
            if cell.row == 1:
                cell.font = Font(bold=True)

    wb.save(file_path)
    print(f"[DEBUG] Событие записано в Excel: {file_path}")
